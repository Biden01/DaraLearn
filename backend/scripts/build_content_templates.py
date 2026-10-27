"""Собирает Excel-шаблоны для команды контента и учителя из учебных материалов в data.py (БД не нужна).

Запуск из папки backend:  python scripts/build_content_templates.py
Результат: ../content/templates/skills_map_template.xlsx и tasks_template.xlsx
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.modules.content.admin import COLUMNS, MIN_TASKS, SKILL_COLUMNS
from app.modules.content.data import COURSES

OUT = Path(__file__).resolve().parent.parent.parent / "content" / "templates"
HEAD = PatternFill("solid", fgColor="DDE7F5")
WRAP = Alignment(wrap_text=True, vertical="top")
PURPOSE_KEY = {"d": "diagnostic", "p": "practice", "v": "verification"}


def _head(ws, widths: list[int]) -> None:
    for c, w in zip(ws[1], widths):
        c.font, c.fill = Font(bold=True), HEAD
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"


def _text(ws, lines: list[str]) -> None:
    ws.column_dimensions["A"].width = 120
    for line in lines:
        ws.append([line])
        ws.cell(ws.max_row, 1).alignment = WRAP
    ws["A1"].font = Font(bold=True, size=14)


def _opt(o) -> str:
    return o if isinstance(o, str) else f"{o['kk']} || {o['ru']}"


def skills_map() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "skills"
    ws.append(SKILL_COLUMNS)
    for c in COURSES:
        for l in c["lessons"]:
            ws.append([c["slug"], l["slug"], l["title"]["kk"], l["title"]["ru"], l["minutes"], ", ".join(l["prereqs"]),
                       "\n".join(l["theory"]["kk"]), "\n".join(l["theory"]["ru"]), "", ""])
    _head(ws, [22, 18, 28, 28, 9, 26, 60, 60, 14, 40])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    dv = DataValidation(type="list", formula1='"да,нет"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"I2:I{ws.max_row + 60}")
    guide = wb.create_sheet("Инструкция")
    _text(guide, [
        "Карта навыков: как заполнять",
        "1. Один навык = одна строка. slug — латиница, цифры и дефис (например, sqrt, linear-eq), не меняется после запуска.",
        "2. prereqs — slug'и навыков, без которых этот не освоить, через запятую. Цикла быть не должно (A требует B, B требует A).",
        "3. Теория: одна строка Excel-ячейки = один абзац. Префикс «$ » — формула, «! » — заметка «Запомни», остальное — обычный текст.",
        "4. Теория нужна на двух языках (kk и ru). Tutor объясняет только по ней: чего нет в теории, он объяснять не будет.",
        "5. Учитель проверяет карту: в колонке teacher_ok ставит «да» или «нет», в teacher_comment — что поправить.",
        "6. Готовый файл загружает администратор: POST /api/content/admin/skills/import. Колонки teacher_* система не читает.",
        "7. Цель пилота: 20–40 навыков. Для каждого — минимум "
        f"{MIN_TASKS['diagnostic']} диагностическое, {MIN_TASKS['practice']} практических и {MIN_TASKS['verification']} проверочных задания (см. tasks_template.xlsx).",
    ])
    return wb


def tasks() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "tasks"
    ws.append(COLUMNS)
    _head(ws, [18, 14, 44, 44, 22, 22, 22, 22, 12, 40, 40, 8])
    dv_p = DataValidation(type="list", formula1='"d,p,v"', allow_blank=True)
    dv_a = DataValidation(type="list", formula1='"1,2,3,4"', allow_blank=True)
    ws.add_data_validation(dv_p)
    ws.add_data_validation(dv_a)
    dv_p.add("B2:B500")
    dv_a.add("I2:I500")

    guide = wb.create_sheet("Инструкция")
    _text(guide, [
        "Банк заданий: как заполнять",
        "1. Лист «tasks» — сюда новые задания. Столбец id оставьте пустым: заполненный id обновляет существующее задание.",
        "2. lesson — slug навыка из карты навыков. purpose: d — диагностика, p — практика (подсказки разрешены), v — проверка (без помощи).",
        "3. Одно задание проверяет один навык. Четыре варианта ответа, верный — номер 1–4 в answer. Неверные варианты — типичные ошибки учеников.",
        "4. Текст на двух языках: kk и ru. Вариант на двух языках — в одной ячейке через « || » (kk || ru). Числа и формулы одинаковы — один текст.",
        "5. explanation — короткое объяснение верного ответа. Для p-заданий это основа подсказки «полный разбор»: пишите по шагам.",
        "6. Для каждого навыка нужны разные по формату v-задания: самостоятельная, «новая ситуация» (другая формулировка того же навыка), мини-проверка.",
        "7. Нельзя копировать задания из ЕНТ-сборников без права на использование. Каждое задание проверяет второй человек (лист «На проверку»).",
        f"8. Минимум на навык: d — {MIN_TASKS['diagnostic']}, p — {MIN_TASKS['practice']}, v — {MIN_TASKS['verification']}. Сводка по текущим навыкам — лист «Покрытие».",
        "9. Загрузка: POST /api/content/admin/tasks/import (администратор). Ошибки возвращаются по номерам строк, остальные строки сохраняются.",
        "",
        "Пример строки (не загружайте её):",
        "sqrt | p | Есепте: √49 | Вычислите: √49 | 5 | 6 | 7 | 8 | 3 | √49 = 7, себебі 7² = 49. | √49 = 7, так как 7² = 49. |",
    ])

    cover = wb.create_sheet("Покрытие")
    cover.append(["slug", "diagnostic", "practice", "verification", "нужно d", "нужно p", "нужно v", "не хватает"])
    for c in COURSES:
        for l in c["lessons"]:
            have = {k: sum(1 for t in l["tasks"] if PURPOSE_KEY[t[0]] == k) for k in MIN_TASKS}
            miss = sum(max(0, MIN_TASKS[k] - have[k]) for k in MIN_TASKS)
            cover.append([l["slug"], have["diagnostic"], have["practice"], have["verification"],
                          MIN_TASKS["diagnostic"], MIN_TASKS["practice"], MIN_TASKS["verification"], miss])
    _head(cover, [20, 12, 12, 14, 10, 10, 10, 12])

    review = wb.create_sheet("На проверку")
    review.append([*COLUMNS[:11], "teacher_ok (да/нет)", "teacher_comment"])
    for c in COURSES:
        for l in c["lessons"]:
            for p, q, opts, r, ex in l["tasks"]:
                o = [_opt(x) for x in opts] + [""] * (4 - len(opts))
                review.append([l["slug"], p, q["kk"], q["ru"], *o[:4], r + 1, ex["kk"], ex["ru"], "", ""])
    _head(review, [18, 10, 40, 40, 20, 20, 20, 20, 10, 36, 36, 12, 36])
    return wb


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    skills_map().save(OUT / "skills_map_template.xlsx")
    tasks().save(OUT / "tasks_template.xlsx")
    print("Готово:", *sorted(p.name for p in OUT.iterdir()))
