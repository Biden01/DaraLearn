"""Админка материалов: редактирование уроков и заданий, импорт/экспорт банка заданий в Excel."""
import io
from typing import Literal

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.db import get_db
from app.core.security import require_role
from app.modules.content.models import Course, Lesson, Task
from app.modules.events.models import LearningEvent
from app.modules.school.models import User
from app.modules.school.router import audit

router = APIRouter(prefix="/content/admin", tags=["content-admin"])

PURPOSES = {"d": "diagnostic", "p": "practice", "v": "verification", "x": "archived"}
COLUMNS = ["lesson", "purpose (d/p/v)", "question_kk", "question_ru", "option1", "option2", "option3", "option4",
           "answer (1-4)", "explanation_kk", "explanation_ru", "id"]
SEP = " || "  # вариант на двух языках: «kk || ru»


class Text2(BaseModel):
    kk: str = Field(max_length=4000)
    ru: str = Field(max_length=4000)


class LessonIn(BaseModel):
    title: Text2
    minutes: int = Field(ge=1, le=180)
    theory: dict[str, list[str]]
    examples: list[dict]
    prereqs: list[str] = []


class NewLessonIn(LessonIn):
    slug: str = Field(pattern=r"^[a-z0-9-]{3,80}$")
    course: str


class TaskIn(BaseModel):
    lesson: str
    purpose: Literal["diagnostic", "practice", "verification", "archived"]
    question: Text2
    options: list[str | Text2] = Field(min_length=2, max_length=6)
    answer: int = Field(ge=0, le=5)
    explanation: Text2


def _task_out(t: Task, lesson_slug: str) -> dict:
    return {"id": t.id, "slug": t.slug, "lesson": lesson_slug, "purpose": t.purpose, "question": t.question,
            "options": t.options, "answer": t.answer, "explanation": t.explanation, "source": t.source}


def _lesson_or_404(db: Session, slug: str) -> Lesson:
    l = db.scalar(select(Lesson).where(Lesson.slug == slug))
    if not l:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "lesson_not_found")
    return l


def _check_prereqs(db: Session, slug: str, prereqs: list[str]) -> None:
    known = set(db.scalars(select(Lesson.slug)))
    bad = [p for p in prereqs if p not in known or p == slug]
    if bad:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"bad_prereqs:{','.join(bad)}")


@router.get("/lessons")
def lessons(user: User = Depends(require_role("admin")), db: Session = Depends(get_db)) -> list[dict]:
    counts = dict(db.execute(select(Task.lesson_id, func.count()).where(Task.purpose != "archived").group_by(Task.lesson_id)).all())
    rows = db.execute(select(Lesson, Course).join(Course, Course.id == Lesson.course_id).order_by(Course.order, Lesson.order)).all()
    return [{"slug": l.slug, "title": l.title, "course": c.slug, "course_title": c.title, "prereqs": l.prereqs,
             "tasks": counts.get(l.id, 0), "source": l.source} for l, c in rows]


@router.get("/lessons/{slug}")
def lesson(slug: str, user: User = Depends(require_role("admin")), db: Session = Depends(get_db)) -> dict:
    l = _lesson_or_404(db, slug)
    tasks = db.scalars(select(Task).where(Task.lesson_id == l.id).order_by(Task.id))
    return {"slug": l.slug, "title": l.title, "minutes": l.minutes, "theory": l.theory, "examples": l.examples,
            "prereqs": l.prereqs, "source": l.source, "tasks": [_task_out(t, l.slug) for t in tasks]}


@router.put("/lessons/{slug}")
def update_lesson(slug: str, payload: LessonIn, user: User = Depends(require_role("admin")), db: Session = Depends(get_db)) -> dict:
    l = _lesson_or_404(db, slug)
    _check_prereqs(db, slug, payload.prereqs)
    l.title, l.minutes, l.theory, l.examples = payload.title.model_dump(), payload.minutes, payload.theory, payload.examples
    l.prereqs, l.source = payload.prereqs, "admin"
    audit(db, user, "lesson.update", "lesson", l.id, slug=slug)
    db.commit()
    return {"slug": l.slug}


@router.post("/lessons", status_code=status.HTTP_201_CREATED)
def create_lesson(payload: NewLessonIn, user: User = Depends(require_role("admin")), db: Session = Depends(get_db)) -> dict:
    course = db.scalar(select(Course).where(Course.slug == payload.course))
    if not course:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "course_not_found")
    if db.scalar(select(Lesson).where(Lesson.slug == payload.slug)):
        raise HTTPException(status.HTTP_409_CONFLICT, "slug_taken")
    _check_prereqs(db, payload.slug, payload.prereqs)
    order = (db.scalar(select(func.max(Lesson.order)).where(Lesson.course_id == course.id)) or 0) + 1
    l = Lesson(slug=payload.slug, course_id=course.id, title=payload.title.model_dump(), minutes=payload.minutes,
               theory=payload.theory, examples=payload.examples, prereqs=payload.prereqs, order=order, source="admin")
    db.add(l)
    db.flush()
    audit(db, user, "lesson.create", "lesson", l.id, slug=l.slug)
    db.commit()
    return {"slug": l.slug}


def _opt(o):
    return o.model_dump() if isinstance(o, Text2) else o


@router.post("/tasks", status_code=status.HTTP_201_CREATED)
def create_task(payload: TaskIn, user: User = Depends(require_role("admin")), db: Session = Depends(get_db)) -> dict:
    l = _lesson_or_404(db, payload.lesson)
    if payload.answer >= len(payload.options):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "bad_answer")
    n = db.scalar(select(func.count()).select_from(Task).where(Task.lesson_id == l.id)) or 0
    t = Task(slug=f"{l.slug}-a{n + 1}-{user.id}", lesson_id=l.id, purpose=payload.purpose, question=payload.question.model_dump(),
             options=[_opt(o) for o in payload.options], answer=payload.answer, explanation=payload.explanation.model_dump(),
             source="admin")
    db.add(t)
    db.flush()
    audit(db, user, "task.create", "task", t.id)
    db.commit()
    return _task_out(t, l.slug)


@router.put("/tasks/{task_id}")
def update_task(task_id: int, payload: TaskIn, user: User = Depends(require_role("admin")), db: Session = Depends(get_db)) -> dict:
    t = db.get(Task, task_id)
    if not t:
        raise HTTPException(status.HTTP_404_NOT_FOUND)
    if payload.answer >= len(payload.options):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "bad_answer")
    l = _lesson_or_404(db, payload.lesson)
    t.lesson_id, t.purpose, t.question = l.id, payload.purpose, payload.question.model_dump()
    t.options, t.answer, t.explanation, t.source = [_opt(o) for o in payload.options], payload.answer, payload.explanation.model_dump(), "admin"
    audit(db, user, "task.update", "task", t.id)
    db.commit()
    return _task_out(t, l.slug)


@router.delete("/tasks/{task_id}", status_code=status.HTTP_204_NO_CONTENT)
def archive_task(task_id: int, user: User = Depends(require_role("admin")), db: Session = Depends(get_db)) -> None:
    """Задание с ответами учеников не удаляется, а уходит в архив — история событий сохраняется."""
    t = db.get(Task, task_id)
    if not t:
        raise HTTPException(status.HTTP_404_NOT_FOUND)
    used = db.scalar(select(func.count()).select_from(LearningEvent).where(LearningEvent.task_id == t.id))
    if used:
        t.purpose, t.source = "archived", "admin"
    else:
        db.delete(t)
    audit(db, user, "task.archive", "task", task_id)
    db.commit()


def _cell_text(o) -> str:
    return o if isinstance(o, str) else f"{o.get('kk', '')}{SEP}{o.get('ru', '')}"


@router.get("/tasks.xlsx")
def export_tasks(user: User = Depends(require_role("admin")), db: Session = Depends(get_db)) -> Response:
    wb = Workbook()
    ws = wb.active
    ws.title = "tasks"
    ws.append(COLUMNS)
    rev = {v: k for k, v in PURPOSES.items()}
    rows = db.execute(select(Task, Lesson.slug).join(Lesson, Lesson.id == Task.lesson_id).order_by(Lesson.id, Task.id)).all()
    for t, slug in rows:
        opts = [_cell_text(o) for o in t.options] + [""] * (4 - len(t.options))
        ws.append([slug, rev[t.purpose], t.question.get("kk", ""), t.question.get("ru", ""), *opts[:4], t.answer + 1,
                   t.explanation.get("kk", ""), t.explanation.get("ru", ""), t.id])
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="daralearn-tasks.xlsx"'})


@router.post("/tasks/import")
async def import_tasks(file: UploadFile = File(...), user: User = Depends(require_role("admin")), db: Session = Depends(get_db)) -> dict:
    """Строка с id — обновление задания, без id — новое. Ошибки возвращаются по номерам строк, остальное сохраняется."""
    data = await file.read(5 * 1024 * 1024)
    try:
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
    except Exception:  # noqa: BLE001
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "bad_file")
    lessons = {l.slug: l for l in db.scalars(select(Lesson))}
    created = updated = 0
    errors = []
    for i, r in enumerate(rows, start=2):
        r = [("" if v is None else str(v)).strip() for v in (list(r) + [None] * 12)[:12]]
        if not any(r[:11]):
            continue
        slug, purpose, qkk, qru, *opts = r[:8]
        ans, ekk, eru, tid = r[8], r[9], r[10], r[11]
        opts = [o for o in opts if o]
        problem = (None if slug in lessons else "lesson") or (None if purpose in PURPOSES else "purpose") \
            or (None if qkk and qru else "question") or (None if len(opts) >= 2 else "options") \
            or (None if ans.isdigit() and 1 <= int(float(ans)) <= len(opts) else "answer")
        if problem:
            errors.append({"row": i, "field": problem})
            continue
        options = [{"kk": o.split(SEP)[0].strip(), "ru": o.split(SEP)[1].strip()} if SEP in o else o for o in opts]
        fields = dict(lesson_id=lessons[slug].id, purpose=PURPOSES[purpose], question={"kk": qkk, "ru": qru},
                      options=options, answer=int(float(ans)) - 1, explanation={"kk": ekk, "ru": eru or ekk}, source="admin")
        t = db.get(Task, int(float(tid))) if tid.replace(".", "").isdigit() else None
        if t:
            for k, v in fields.items():
                setattr(t, k, v)
            updated += 1
        else:
            n = db.scalar(select(func.count()).select_from(Task).where(Task.lesson_id == fields["lesson_id"])) or 0
            db.add(Task(slug=f"{slug}-i{n + 1}-{i}", **fields))
            db.flush()
            created += 1
    audit(db, user, "task.import", "task", None, created=created, updated=updated, errors=len(errors))
    db.commit()
    return {"created": created, "updated": updated, "errors": errors}


# ---- карта навыков: Excel для проверки учителем ----

SKILL_COLUMNS = ["course", "slug", "title_kk", "title_ru", "minutes", "prereqs (через запятую)", "theory_kk", "theory_ru",
                 "teacher_ok (да/нет)", "teacher_comment"]
# минимум заданий на навык: хватает на диагностику, закрепление (4) и проверку (3 без повторов)
MIN_TASKS = {"diagnostic": 1, "practice": 4, "verification": 3}


def _has_cycle(graph: dict[str, list[str]]) -> list[str] | None:
    state: dict[str, int] = {}
    stack: list[str] = []

    def visit(n: str) -> list[str] | None:
        state[n] = 1
        stack.append(n)
        for m in graph.get(n, []):
            if state.get(m) == 1:
                return stack[stack.index(m):] + [m]
            if m not in state and (c := visit(m)):
                return c
        stack.pop()
        state[n] = 2
        return None

    for n in graph:
        if n not in state and (c := visit(n)):
            return c
    return None


@router.get("/skills.xlsx")
def export_skills(user: User = Depends(require_role("admin")), db: Session = Depends(get_db)) -> Response:
    wb = Workbook()
    ws = wb.active
    ws.title = "skills"
    ws.append(SKILL_COLUMNS)
    for l, c in db.execute(select(Lesson, Course).join(Course, Course.id == Lesson.course_id).order_by(Course.order, Lesson.order)).all():
        ws.append([c.slug, l.slug, l.title.get("kk", ""), l.title.get("ru", ""), l.minutes, ", ".join(l.prereqs),
                   "\n".join(l.theory.get("kk", [])), "\n".join(l.theory.get("ru", [])), "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="daralearn-skills.xlsx"'})


@router.post("/skills/import")
async def import_skills(file: UploadFile = File(...), user: User = Depends(require_role("admin")),
                        db: Session = Depends(get_db)) -> dict:
    """Карта навыков из Excel. Строка с существующим slug обновляет название, время, prerequisite и теорию
    (примеры не трогает), новый slug создаёт навык. Цикл в prerequisite отклоняет весь файл."""
    data = await file.read(5 * 1024 * 1024)
    try:
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
    except Exception:  # noqa: BLE001
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "bad_file")
    courses = {c.slug: c for c in db.scalars(select(Course))}
    existing = {l.slug: l for l in db.scalars(select(Lesson))}
    parsed, errors = [], []
    import re
    for i, r in enumerate(rows, start=2):
        r = [("" if v is None else str(v)).strip() for v in (list(r) + [None] * 8)[:8]]
        if not any(r):
            continue
        course, slug, tkk, tru, mins, prereqs, thkk, thru = r
        prereq_list = [p.strip() for p in prereqs.split(",") if p.strip()]
        problem = (None if course in courses else "course") or (None if re.fullmatch(r"[a-z0-9-]{3,80}", slug) else "slug") \
            or (None if tkk and tru else "title") or (None if mins.replace(".0", "").isdigit() and 1 <= int(float(mins)) <= 180 else "minutes") \
            or (None if thkk and thru else "theory")
        if problem:
            errors.append({"row": i, "field": problem})
            continue
        parsed.append(dict(row=i, course=course, slug=slug, title={"kk": tkk, "ru": tru}, minutes=int(float(mins)),
                           prereqs=prereq_list, theory={"kk": [x for x in thkk.splitlines() if x.strip()],
                                                        "ru": [x for x in thru.splitlines() if x.strip()]}))
    known = set(existing) | {p["slug"] for p in parsed}
    for p in parsed:
        bad = [x for x in p["prereqs"] if x not in known or x == p["slug"]]
        if bad:
            errors.append({"row": p["row"], "field": "prereqs", "values": bad})
    if errors:
        return {"created": 0, "updated": 0, "errors": errors}
    graph = {s: list(l.prereqs) for s, l in existing.items()} | {p["slug"]: p["prereqs"] for p in parsed}
    if cycle := _has_cycle(graph):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"prereq_cycle:{'>'.join(cycle)}")
    created = updated = 0
    for p in parsed:
        l = existing.get(p["slug"])
        if l:
            l.title, l.minutes, l.prereqs, l.theory, l.source = p["title"], p["minutes"], p["prereqs"], p["theory"], "admin"
            updated += 1
        else:
            order = (db.scalar(select(func.max(Lesson.order)).where(Lesson.course_id == courses[p["course"]].id)) or 0) + 1
            db.add(Lesson(slug=p["slug"], course_id=courses[p["course"]].id, title=p["title"], minutes=p["minutes"],
                          theory=p["theory"], examples=[], prereqs=p["prereqs"], order=order, source="admin"))
            created += 1
    audit(db, user, "skills.import", "lesson", None, created=created, updated=updated)
    db.commit()
    return {"created": created, "updated": updated, "errors": []}


@router.get("/coverage")
def coverage(user: User = Depends(require_role("admin")), db: Session = Depends(get_db)) -> dict:
    """Чего не хватает для пилота: по каждому навыку задания по назначению и наличие теории."""
    counts: dict[int, dict[str, int]] = {}
    for lid, purpose, n in db.execute(select(Task.lesson_id, Task.purpose, func.count()).group_by(Task.lesson_id, Task.purpose)):
        counts.setdefault(lid, {})[purpose] = n
    rows = db.execute(select(Lesson, Course).join(Course, Course.id == Lesson.course_id).order_by(Course.order, Lesson.order)).all()
    skills = []
    for l, c in rows:
        have = counts.get(l.id, {})
        missing = {p: need - have.get(p, 0) for p, need in MIN_TASKS.items() if have.get(p, 0) < need}
        theory = all(bool(l.theory.get(lang)) for lang in ("kk", "ru"))
        skills.append({"slug": l.slug, "course": c.slug, "have": {p: have.get(p, 0) for p in MIN_TASKS}, "missing": missing,
                       "theory": theory, "ready": not missing and theory})
    return {"skills": skills, "total": len(skills), "ready": sum(s["ready"] for s in skills),
            "tasks_total": sum(sum(s["have"].values()) for s in skills), "tasks_needed": len(skills) * sum(MIN_TASKS.values()),
            "min_per_skill": MIN_TASKS}
