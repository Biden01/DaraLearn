"""Импорт учеников из Excel/CSV: логины латиницей и временные пароли."""
import csv
import io
import re
import secrets

from openpyxl import Workbook, load_workbook

TRANSLIT = dict(zip("абвгдеёжзийклмнопрстуфхцчшщъыьэюяәғқңөұүһі",
                    ["a", "b", "v", "g", "d", "e", "e", "zh", "z", "i", "i", "k", "l", "m", "n", "o", "p", "r", "s", "t",
                     "u", "f", "kh", "ts", "ch", "sh", "sh", "", "y", "", "e", "iu", "ia", "a", "g", "q", "n", "o", "u",
                     "u", "h", "i"]))


def translit(s: str) -> str:
    out = "".join(TRANSLIT.get(ch, ch) for ch in s.lower())
    return re.sub(r"[^a-z0-9]", "", out)


def base_login(full_name: str) -> str:
    parts = [p for p in re.split(r"\s+", full_name.strip()) if p]
    if not parts:
        return "student"
    # «Фамилия Имя» → имя.ф
    first = translit(parts[1] if len(parts) > 1 else parts[0]) or "student"
    initial = translit(parts[0])[:1] if len(parts) > 1 else ""
    return f"{first}.{initial}" if initial else first


def temp_password() -> str:
    return "".join(secrets.choice("abcdefghjkmnpqrstuvwxyz23456789") for _ in range(8))


def read_rows(filename: str, data: bytes, keep_header: bool = False) -> list[list[str]]:
    """Строки таблицы (по умолчанию без заголовка). Поддерживаются .xlsx и .csv (; или ,)."""
    if filename.lower().endswith(".xlsx"):
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = [[("" if v is None else str(v)).strip() for v in r] for r in ws.iter_rows(values_only=True)]
    else:
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=";,\t") if text.strip() else csv.excel
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    rows = [r for r in rows if any(r)]
    if not keep_header and rows and not any(ch.isdigit() for ch in "".join(rows[0])) and len(rows) > 1:
        header = " ".join(rows[0]).lower()
        if any(w in header for w in ("фио", "name", "аты", "имя", "логин", "login", "задание", "task")):
            rows = rows[1:]
    return rows


def credentials_xlsx(rows: list[tuple[str, str, str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "DaraLearn"
    ws.append(["ФИО / Аты-жөні", "Логин", "Временный пароль / Уақытша құпиясөз"])
    for r in rows:
        ws.append(list(r))
    for col, w in zip("ABC", (34, 20, 26)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
