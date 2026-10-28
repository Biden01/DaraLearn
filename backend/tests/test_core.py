import io

from openpyxl import Workbook, load_workbook

from tests.conftest import diagnostic, register, uid


def test_auth_and_roles(admin):
    me = admin.ok("GET", "/api/auth/me")
    assert me["role"] == "admin"
    s = register()
    assert s.ok("GET", "/api/auth/me")["role"] == "student"
    assert s.get("/api/admin/metrics").status_code == 403
    assert s.get("/api/school/classes").status_code == 403
    s.ok("POST", "/api/auth/password", 204, json={"old": "secret123", "new": "newsecret1"})
    assert s.post("/api/auth/password", json={"old": "wrong", "new": "newsecret2"}).status_code == 400


def test_diagnostic_persists_root_cause():
    s = register()
    r = diagnostic(s, wrong={"linear-eq", "qe-def", "discriminant"})
    d = s.ok("GET", "/api/decisions/me")
    by = {x["lesson"]: x for x in d["recommendations"]}
    assert by["linear-eq"]["action"] == "lesson"
    # каскад: дискриминант → корень цепочки — линейные уравнения
    assert by["discriminant"]["action"] == "prereq" and by["discriminant"]["target"] == "linear-eq"
    assert all(x["status"] == "new" for x in d["recommendations"])
    st = s.ok("GET", "/api/state/me")
    assert st["progress"]["mastered"] == 0  # один верный ответ — только «предварительно»
    assert {x["status"] for x in st["lessons"]} == {"likely", "gap"}
    assert r["recommendations"]


def test_school_codes_import_and_access(school):
    t, cls = school["teacher"], school["class"]
    assert t.ok("GET", "/api/auth/me")["role"] == "teacher"
    student = register(code=cls["code"])
    assert [c["name"] for c in student.ok("GET", "/api/auth/me/classes")] == ["8А"]
    # импорт списка → Excel с логинами
    wb = Workbook()
    ws = wb.active
    ws.append(["ФИО"])
    ws.append(["Нұрланов Әлихан"])
    ws.append(["Иванова Мария"])
    buf = io.BytesIO()
    wb.save(buf)
    r = t.ok("POST", f"/api/school/classes/{cls['id']}/import",
             files={"file": ("list.xlsx", buf.getvalue(), "application/octet-stream")})
    rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2 and rows[0][1].startswith("alikhan.n") and rows[1][1].startswith("mariia.i")
    from tests.conftest import client
    kid = client()
    me = kid.ok("POST", "/api/auth/login", json={"login": rows[0][1], "password": rows[0][2]})
    assert me["must_change_password"] is True
    assert len(t.ok("GET", f"/api/school/classes/{cls['id']}/students")) == 3
    # другой учитель не видит чужой класс
    other = register(code=school["admin"].ok("POST", f"/api/school/schools/{school['school']['id']}/invites", 201,
                                              json={"role": "teacher"})["code"])
    assert other.get(f"/api/state/classes/{cls['id']}").status_code == 404
    assert other.get(f"/api/state/students/{me['id']}").status_code == 404
    # куратор школы видит
    assert school["curator"].ok("GET", f"/api/state/classes/{cls['id']}")["class"]["name"] == "8А"


def test_teacher_feed_decisions_and_group(school):
    t, cls = school["teacher"], school["class"]
    kids = [register(code=cls["code"]) for _ in range(3)]
    for k in kids:
        diagnostic(k, wrong={"sqrt"})
    feed = t.ok("GET", f"/api/decisions/feed?class_id={cls['id']}")
    group = [i for i in feed["items"] if i["action"] == "group"]
    assert group and group[0]["lesson"] == "sqrt"
    personal = [i for i in feed["items"] if i.get("student") and i["lesson"] == "sqrt"]
    assert len(personal) == 3
    rec = personal[0]
    t.ok("PATCH", f"/api/decisions/{rec['id']}", json={"status": "changed", "action": "tutor", "comment": "Разберём с Tutor"})
    t.ok("PATCH", f"/api/decisions/{personal[1]['id']}", json={"status": "rejected"})
    kid = next(k for k in kids if k.ok("GET", "/api/auth/me")["id"] == rec["user_id"])
    mine = kid.ok("GET", "/api/decisions/me")["recommendations"]  # повторная синхронизация
    sq = [r for r in mine if r["lesson"] == "sqrt"]
    assert sq and sq[0]["action"] == "tutor" and sq[0]["status"] == "changed" and sq[0]["comment"] == "Разберём с Tutor"
    rejected_kid = next(k for k in kids if k.ok("GET", "/api/auth/me")["id"] == personal[1]["user_id"])
    assert not [r for r in rejected_kid.ok("GET", "/api/decisions/me")["recommendations"] if r["lesson"] == "sqrt"]
    heat = t.ok("GET", f"/api/state/classes/{cls['id']}")
    sqrt_col = next(l for l in heat["lessons"] if l["slug"] == "sqrt")
    assert sqrt_col["gaps"] == 3


def _make_due(db, user_id):
    from datetime import datetime, timedelta, timezone

    from app.modules.verification.models import Check
    for c in db.query(Check).filter(Check.user_id == user_id, Check.status == "pending"):
        c.due_at = datetime.now(timezone.utc) - timedelta(minutes=1)
    db.commit()


def test_practice_check_outcome(db):
    s = register()
    diagnostic(s, wrong={"sqrt"})
    lesson = s.ok("GET", "/api/content/lessons/sqrt")
    from app.modules.content.models import Task
    t = db.get(Task, lesson["practice"][0]["id"])
    assert s.ok("POST", "/api/events/practice", json={"task_id": t.id, "choice": t.answer})["correct"]
    checks = s.ok("GET", "/api/verification/checks")
    assert len(checks) == 1 and checks[0]["due"] is False
    me = s.ok("GET", "/api/auth/me")
    _make_due(db, me["id"])
    assert s.ok("GET", "/api/state/me")["checks_due"] == 1
    a = s.ok("POST", "/api/verification/attempts", 201, json={"kind": "verification", "check_id": checks[0]["id"]})
    key = {x.id: x.answer for x in db.query(Task).filter(Task.id.in_([x["id"] for x in a["tasks"]]))}
    r = s.ok("POST", f"/api/verification/attempts/{a['id']}/submit", json={"answers": {str(k): v for k, v in key.items()}})
    assert r["outcome"] and r["outcome"]["helped"] is True and r["outcome"]["after"] > r["outcome"]["before"]
    assert s.ok("GET", "/api/verification/checks") == []


def test_tutor_stages_and_teacher_rule(db):
    s = register()
    diagnostic(s, wrong={"sqrt"})
    for _ in range(2):
        sess = s.ok("POST", "/api/tutor/sessions", 201, json={"lesson": "sqrt"})
        s.ok("POST", f"/api/tutor/sessions/{sess['id']}/messages", json={"text": "Помоги"})  # заглушка без ключа
        s.ok("POST", f"/api/tutor/sessions/{sess['id']}/stage", json={"stage": 6})
        chk = s.ok("GET", f"/api/tutor/sessions/{sess['id']}/check")
        assert chk["session"]["stage"] == 7 and len(chk["tasks"]) == 3
        assert s.post(f"/api/tutor/sessions/{sess['id']}/messages", json={"text": "подскажи"}).status_code == 409
        from app.modules.content.models import Task
        wrong = {str(x["id"]): (db.get(Task, x["id"]).answer + 1) % 4 for x in chk["tasks"]}
        fin = s.ok("POST", f"/api/tutor/sessions/{sess['id']}/check", json={"answers": wrong})
        assert fin["session"]["result"] == "failed" and fin["session"]["stage"] == 9
    recs = {r["lesson"]: r for r in s.ok("GET", "/api/decisions/me")["recommendations"]}
    assert recs["sqrt"]["action"] == "teacher"  # 2 помощи без роста → учитель


def test_events_import(school):
    t, cls = school["teacher"], school["class"]
    kid = register(code=cls["code"], email=f"imp{uid()}@test.kz")
    email = kid.ok("GET", "/api/auth/me")["email"]
    csv = f"логин;задание;верно;дата\n{email};sqrt;0;2026-10-01\n{email};sqrt-2;нет;01.10.2026\nnobody;sqrt;1;\n{email};zzz;1;\n"
    r = t.ok("POST", "/api/events/import", files={"file": ("r.csv", csv.encode(), "text/csv")})
    assert r["saved"] == 2 and r["error_count"] == 2
    assert {e["field"] for e in r["errors"]} == {"student", "task"}
    st = {x["slug"]: x for x in kid.ok("GET", "/api/state/me")["lessons"]}
    assert st["sqrt"]["status"] == "gap" and st["sqrt"]["n_indep"] == 2


def test_adviser_case(school, db):
    from datetime import datetime, timedelta, timezone

    from app.modules.verification.models import Check
    t, cls = school["teacher"], school["class"]
    kid = register(code=cls["code"])
    diagnostic(kid, wrong={"sqrt", "brackets"})
    t.ok("POST", f"/api/verification/classes/{cls['id']}/assign", 201, json={"lesson": "sqrt"})
    me = kid.ok("GET", "/api/auth/me")
    c = db.query(Check).filter(Check.user_id == me["id"]).first()
    c.due_at = datetime.now(timezone.utc) - timedelta(days=3)
    db.add(Check(user_id=me["id"], lesson_id=c.lesson_id + 1, reason="lesson", before_proof=0,
                 due_at=datetime.now(timezone.utc) - timedelta(days=2)))
    db.commit()
    queue = school["curator"].ok("GET", "/api/verification/adviser")
    case = next(x for x in queue if x["student"]["id"] == me["id"])
    assert case["plan_rate"] < 0.5 and case["student"]["class"] == "8А"
    upd = school["curator"].ok("PATCH", f"/api/verification/adviser/{case['id']}",
                               json={"status": "watching", "next_step": "Разговор с родителями", "watch_until": "2026-11-10"})
    assert upd["status"] == "watching"


def test_content_admin_roundtrip(admin):
    lessons = admin.ok("GET", "/api/content/admin/lessons")
    assert len(lessons) >= 12
    x = admin.get("/api/content/admin/tasks.xlsx")
    wb = load_workbook(io.BytesIO(x.content))
    ws = wb.active
    first = [c.value for c in ws[2]]
    ws.append(["sqrt", "p", "Жаңа сұрақ", "Новый вопрос", "1", "2", "3", "4", 2, "т", "о", None])
    ws.append(["nope", "p", "q", "q", "1", "2", None, None, 1, "", "", None])
    buf = io.BytesIO()
    wb.save(buf)
    r = admin.ok("POST", "/api/content/admin/tasks/import", files={"file": ("t.xlsx", buf.getvalue(), "application/octet-stream")})
    assert r["created"] == 1 and r["errors"] == [{"row": ws.max_row, "field": "lesson"}] and r["updated"] >= 1
    full = admin.ok("GET", "/api/content/admin/lessons/sqrt")
    assert full["tasks"][0]["source"] == "admin" or first[11]
    new = [t for t in full["tasks"] if t["question"]["ru"] == "Новый вопрос"][0]
    admin.ok("DELETE", f"/api/content/admin/tasks/{new['id']}", 204)
    full["title"]["ru"] = "Квадратный корень (ред.)"
    admin.ok("PUT", "/api/content/admin/lessons/sqrt", json={k: full[k] for k in ("title", "minutes", "theory", "examples", "prereqs")})
    assert admin.ok("GET", "/api/content/admin/lessons/sqrt")["source"] == "admin"


def test_admin_metrics(admin):
    m = admin.ok("GET", "/api/admin/metrics")
    assert "independent_progress" in m and "accepted_share" in m["recommendations"]
    assert admin.ok("GET", "/api/admin/ai")["models"] is not None
    assert isinstance(admin.ok("GET", "/api/admin/audit"), list)
