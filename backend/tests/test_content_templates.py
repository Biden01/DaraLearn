import io

from openpyxl import Workbook, load_workbook

import pytest

from tests.conftest import uid


@pytest.fixture(autouse=True)
def _cleanup_lessons(db):
    """Тесты создают навыки без заданий; в общей базе они ломают другие тесты, поэтому убираем за собой."""
    yield
    from app.modules.content.models import Lesson

    db.rollback()
    db.query(Lesson).filter(Lesson.slug.like("skill-%") | Lesson.slug.like("cyc-%")).delete(synchronize_session=False)
    db.commit()


def _xlsx(rows):
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(admin, rows, expect=200):
    return admin.ok("POST", "/api/content/admin/skills/import", expect,
                    files={"file": ("skills.xlsx", _xlsx(rows), "application/octet-stream")})


HEAD = ["course", "slug", "title_kk", "title_ru", "minutes", "prereqs", "theory_kk", "theory_ru"]


def test_skills_export_import_roundtrip(admin):
    r = admin.ok("GET", "/api/content/admin/skills.xlsx")
    rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(min_row=2, values_only=True))
    assert any(x[1] == "sqrt" for x in rows)
    # новый навык с prerequisite на существующий; существующий не ломается
    slug = f"skill-{uid()}"
    out = _upload(admin, [HEAD, ["quadratic-equations", slug, "Атауы", "Название", 15, "sqrt, brackets", "Теория\n$ x²", "Теория ru"]])
    assert out["created"] == 1 and not out["errors"]
    lesson = admin.ok("GET", f"/api/content/admin/lessons/{slug}")
    assert lesson["prereqs"] == ["sqrt", "brackets"] and lesson["theory"]["kk"] == ["Теория", "$ x²"]
    assert _upload(admin, [HEAD, ["quadratic-equations", slug, "Жаңа", "Новое", 20, "sqrt", "A", "B"]])["updated"] == 1


def test_skills_import_rejects_bad_rows_and_cycles(admin):
    a, b = f"cyc-a-{uid()}", f"cyc-b-{uid()}"
    bad = _upload(admin, [HEAD, ["nope", a, "x", "y", 10, "", "t", "t"], ["quadratic-equations", b, "x", "y", 10, "ghost", "t", "t"]])
    assert {e["field"] for e in bad["errors"]} == {"course", "prereqs"} and bad["created"] == 0
    r = admin.post("/api/content/admin/skills/import", files={"file": ("s.xlsx", _xlsx([
        HEAD, ["quadratic-equations", a, "x", "y", 10, b, "t", "t"], ["quadratic-equations", b, "x", "y", 10, a, "t", "t"]]),
        "application/octet-stream")})
    assert r.status_code == 400 and r.json()["detail"].startswith("prereq_cycle")
    assert admin.get(f"/api/content/admin/lessons/{a}").status_code == 404  # ничего не сохранилось


def test_coverage_reports_what_is_missing(admin):
    c = admin.ok("GET", "/api/content/admin/coverage")
    sqrt = next(s for s in c["skills"] if s["slug"] == "sqrt")
    assert set(sqrt["have"]) == {"diagnostic", "practice", "verification"} and c["tasks_needed"] == c["total"] * 8
    assert c["ready"] <= c["total"]


def test_generated_templates_are_importable(admin):
    import sys
    from pathlib import Path

    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
    import build_content_templates as b

    buf = io.BytesIO()
    b.skills_map().save(buf)
    out = admin.ok("POST", "/api/content/admin/skills/import",
                   files={"file": ("s.xlsx", buf.getvalue(), "application/octet-stream")})
    assert not out["errors"] and out["updated"] >= 12
    buf = io.BytesIO()
    b.tasks().save(buf)  # пустой лист tasks: загрузка ничего не создаёт и не ломается
    out = admin.ok("POST", "/api/content/admin/tasks/import",
                   files={"file": ("t.xlsx", buf.getvalue(), "application/octet-stream")})
    assert out == {"created": 0, "updated": 0, "errors": []}
